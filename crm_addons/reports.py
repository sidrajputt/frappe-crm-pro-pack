"""The Sales Dashboard: calls, follow-ups, leads and meetings, per day and per person.

Who sees what follows the roles. Sales Managers, System Managers and Administrator see everyone
(and can look at one person); a Sales User only ever sees their own numbers, whatever is asked.
"""

import frappe
from frappe import _
from frappe.utils import add_days, cint, date_diff, getdate, now_datetime, nowdate

from crm_addons import scoring, stale
from crm_addons.utils import get_reference_info, get_settings, is_manager, user_details

ROLES = ("Sales User", "Sales Manager", "System Manager")
COUNTS = ("calls", "connected", "not_picked", "not_connected", "attempt_sum", "follow_ups", "leads", "meetings", "held", "converted")


def _period(from_date, to_date):
	start = getdate(from_date or add_days(nowdate(), -6))
	end = getdate(to_date or nowdate())
	if end < start:
		start, end = end, start
	if date_diff(end, start) > 366:
		start = add_days(end, -366)
	return start, end


def _who(user):
	"""The one user to report on, or None for everybody (managers only)."""
	if not is_manager():
		return frappe.session.user
	return user or None


def _metrics(start, end, who):
	"""Per-person counts for [start, end). ``end`` is exclusive."""
	params = {"start": start, "end": end, "who": who}

	def cond(column):
		return f" and {column} = %(who)s" if who else ""

	people = {}

	def row(person):
		return people.setdefault(person or "", dict.fromkeys(COUNTS, 0))

	for r in frappe.db.sql(
		f"""select f.owner as person, count(*) as follow_ups,
			sum(f.mode = 'Call') as calls,
			sum(f.mode = 'Call' and f.call_status = 'Connected') as connected,
			sum(f.mode = 'Call' and f.outcome = 'Did Not Pick') as not_picked,
			sum(f.mode = 'Call' and f.outcome = 'Did Not Connect') as not_connected,
			sum(case when f.mode = 'Call' and f.call_status = 'Connected' then f.attempt_no else 0 end) as attempt_sum
		from `tabCRM Follow Up` f
		where f.followed_up_on >= %(start)s and f.followed_up_on < %(end)s {cond('f.owner')}
		group by f.owner""",
		params,
		as_dict=True,
	):
		row(r.person).update({k: cint(r[k]) for k in ("follow_ups", "calls", "connected", "not_picked", "not_connected", "attempt_sum")})

	for r in frappe.db.sql(
		f"""select l.lead_owner as person, count(*) as leads from `tabCRM Lead` l
		where l.creation >= %(start)s and l.creation < %(end)s {cond('l.lead_owner')} group by l.lead_owner""",
		params,
		as_dict=True,
	):
		row(r.person)["leads"] = cint(r.leads)

	for r in frappe.db.sql(
		f"""select m.organizer as person, count(*) as meetings, sum(m.meeting_outcome = 'Held') as held
		from `tabCRM Meeting` m
		where m.status in ('Scheduled', 'Completed') and m.starts_on >= %(start)s and m.starts_on < %(end)s
		{cond('m.organizer')} group by m.organizer""",
		params,
		as_dict=True,
	):
		row(r.person).update({"meetings": cint(r.meetings), "held": cint(r.held)})

	if frappe.db.table_exists("CRM Deal"):
		for r in frappe.db.sql(
			f"""select d.deal_owner as person, count(*) as converted from `tabCRM Deal` d
			where d.lead is not null and d.lead != '' and d.creation >= %(start)s and d.creation < %(end)s
			{cond('d.deal_owner')} group by d.deal_owner""",
			params,
			as_dict=True,
		):
			row(r.person)["converted"] = cint(r.converted)
	return people


def _pending(who):
	"""Open follow-ups per person: due today, and overdue (before today)."""
	today = getdate(nowdate())
	params = {"today": today, "tomorrow": add_days(today, 1), "who": who}
	people = {}
	who_sql = " and f.assigned_to = %(who)s" if who else ""
	for r in frappe.db.sql(
		f"""select f.assigned_to as person,
			sum(f.next_follow_up_on < %(today)s) as overdue,
			sum(f.next_follow_up_on >= %(today)s and f.next_follow_up_on < %(tomorrow)s) as due_today
		from `tabCRM Follow Up` f
		where f.next_closed = 0 and f.next_follow_up_on is not null
		{who_sql} group by f.assigned_to""",
		params,
		as_dict=True,
	):
		people[r.person or ""] = {"overdue": cint(r.overdue), "due_today": cint(r.due_today)}
	return people


def _sum(people):
	total = dict.fromkeys(COUNTS, 0)
	for values in people.values():
		for k in COUNTS:
			total[k] += values[k]
	return total


def _totals(counts):
	calls, connected = counts["calls"], counts["connected"]
	return {
		**{k: counts[k] for k in ("calls", "connected", "not_picked", "not_connected", "follow_ups", "leads", "meetings", "held", "converted")},
		"connect_rate": round(connected / calls * 100, 1) if calls else 0,
		"avg_attempts": round(counts["attempt_sum"] / connected, 1) if connected else 0,
	}


def _team(who):
	"""Everybody who can be reported on: the sales team, or just the one user."""
	if who:
		return {who}
	roles = frappe.get_all("Has Role", filters={"role": ["in", ["Sales User", "Sales Manager"]], "parenttype": "User"}, pluck="parent")
	return {
		u
		for u in set(roles)
		if u not in ("Administrator", "Guest") and frappe.db.get_value("User", u, "enabled")
	}


def _title(reference_docname):
	info = get_reference_info("CRM Lead", reference_docname)
	return info["title"], info["url"]


def _build_report(from_date=None, to_date=None, user=None):
	start, end = _period(from_date, to_date)
	who = _who(user)
	manager = is_manager()
	params = {"start": start, "end": add_days(end, 1), "who": who}
	span = date_diff(end, start) + 1

	current = _metrics(start, add_days(end, 1), who)
	previous = _metrics(add_days(start, -span), start, who)
	today = getdate(nowdate())
	todays = _metrics(today, add_days(today, 1), who)
	pending = _pending(who)

	def cond(column):
		return f" and {column} = %(who)s" if who else ""

	# ---- per day and person; the page adds them up as it needs
	days = {}
	for r in frappe.db.sql(
		f"""select date(f.followed_up_on) as day, f.owner as person, count(*) as follow_ups,
			sum(f.mode = 'Call') as calls, sum(f.mode = 'Call' and f.call_status = 'Connected') as connected
		from `tabCRM Follow Up` f
		where f.followed_up_on >= %(start)s and f.followed_up_on < %(end)s {cond('f.owner')}
		group by day, f.owner""",
		params,
		as_dict=True,
	):
		days[(str(r.day), r.person or "")] = {"date": str(r.day), "person": r.person or "", "calls": cint(r.calls), "connected": cint(r.connected), "follow_ups": cint(r.follow_ups), "leads": 0}
	for r in frappe.db.sql(
		f"""select date(l.creation) as day, l.lead_owner as person, count(*) as leads from `tabCRM Lead` l
		where l.creation >= %(start)s and l.creation < %(end)s {cond('l.lead_owner')} group by day, l.lead_owner""",
		params,
		as_dict=True,
	):
		key = (str(r.day), r.person or "")
		days.setdefault(key, {"date": str(r.day), "person": r.person or "", "calls": 0, "connected": 0, "follow_ups": 0, "leads": 0})["leads"] = cint(r.leads)

	outcomes = frappe.db.sql(
		f"""select f.outcome as outcome, count(*) as count from `tabCRM Follow Up` f
		where f.mode = 'Call' and f.outcome != '' and f.followed_up_on >= %(start)s and f.followed_up_on < %(end)s {cond('f.owner')}
		group by f.outcome order by count desc""",
		params,
		as_dict=True,
	)

	pipeline = frappe.db.sql(
		f"""select l.status as status, s.color as color, s.position as position, count(*) as count,
			avg(datediff(now(), l.creation)) as age
		from `tabCRM Lead` l left join `tabCRM Lead Status` s on s.name = l.status
		where l.converted = 0 {cond('l.lead_owner')}
		group by l.status, s.color, s.position order by s.position asc, count desc""",
		params,
		as_dict=True,
	)

	bands = {"hot": 0, "warm": 0, "cold": 0}
	if frappe.get_meta("CRM Lead").has_field(scoring.FIELD):
		r = frappe.db.sql(
			f"""select sum(l.addons_score >= %(hot)s) as hot,
				sum(l.addons_score >= %(warm)s and l.addons_score < %(hot)s) as warm,
				sum(l.addons_score < %(warm)s) as cold
			from `tabCRM Lead` l where l.converted = 0 and l.addons_score is not null {cond('l.lead_owner')}""",
			{"hot": scoring.HOT, "warm": scoring.WARM, "who": who},
			as_dict=True,
		)[0]
		bands = {k: cint(r[k]) for k in bands}

	# ---- what became of the leads created in the period
	deal_exists = "exists(select 1 from `tabCRM Deal` d where d.lead = l.name)" if frappe.db.table_exists("CRM Deal") else "0"
	funnel_row = frappe.db.sql(
		f"""select count(*) as created,
			sum(exists(select 1 from `tabCRM Follow Up` x where x.reference_doctype = 'CRM Lead' and x.reference_docname = l.name)) as contacted,
			sum(exists(select 1 from `tabCRM Follow Up` x where x.reference_doctype = 'CRM Lead' and x.reference_docname = l.name
				and x.call_status = 'Connected')) as connected,
			sum(exists(select 1 from `tabCRM Meeting` m where m.reference_doctype = 'CRM Lead' and m.reference_docname = l.name
				and m.status != 'Cancelled')) as meeting,
			sum({deal_exists}) as converted
		from `tabCRM Lead` l where l.creation >= %(start)s and l.creation < %(end)s {cond('l.lead_owner')}""",
		params,
		as_dict=True,
	)[0]
	funnel = [
		{"stage": "Leads added", "count": cint(funnel_row.created)},
		{"stage": "Contacted", "count": cint(funnel_row.contacted)},
		{"stage": "Reached on a call", "count": cint(funnel_row.connected)},
		{"stage": "Meeting booked", "count": cint(funnel_row.meeting)},
		{"stage": "Converted to deal", "count": cint(funnel_row.converted)},
	]

	minutes = frappe.db.sql(
		f"""select avg(greatest(timestampdiff(minute, l.creation, fu.first_at), 0))
		from `tabCRM Lead` l
		join (select x.reference_docname as lead_name, min(x.followed_up_on) as first_at from `tabCRM Follow Up` x
			where x.reference_doctype = 'CRM Lead' group by x.reference_docname) fu on fu.lead_name = l.name
		where l.creation >= %(start)s and l.creation < %(end)s {cond('l.lead_owner')}""",
		params,
	)[0][0]
	speed_hours = round(float(minutes) / 60, 1) if minutes is not None else None

	hours = {r.h: r for r in frappe.db.sql(
		f"""select hour(f.followed_up_on) as h, count(*) as calls, sum(f.call_status = 'Connected') as connected
		from `tabCRM Follow Up` f
		where f.mode = 'Call' and f.followed_up_on >= %(start)s and f.followed_up_on < %(end)s {cond('f.owner')} group by h""",
		params,
		as_dict=True,
	)}
	by_hour = [{"hour": h, "calls": cint(hours[h].calls) if h in hours else 0, "connected": cint(hours[h].connected) if h in hours else 0} for h in range(24)]

	attempts = {r.a: cint(r.n) for r in frappe.db.sql(
		f"""select least(f.attempt_no, 4) as a, count(*) as n from `tabCRM Follow Up` f
		where f.mode = 'Call' and f.call_status = 'Connected' and f.followed_up_on >= %(start)s and f.followed_up_on < %(end)s
		{cond('f.owner')} group by a""",
		params,
		as_dict=True,
	)}
	attempts_rows = [{"label": lbl, "count": attempts.get(i, 0)} for i, lbl in ((1, "1st call"), (2, "2nd call"), (3, "3rd call"), (4, "4th or later"))]

	sources = frappe.db.sql(
		f"""select coalesce(nullif(l.source, ''), 'Unknown') as source, count(*) as leads, sum({deal_exists}) as converted
		from `tabCRM Lead` l where l.creation >= %(start)s and l.creation < %(end)s {cond('l.lead_owner')}
		group by 1 order by leads desc limit 12""",
		params,
		as_dict=True,
	)

	meeting_outcomes = frappe.db.sql(
		f"""select coalesce(nullif(m.meeting_outcome, ''), 'Not recorded') as outcome, count(*) as count from `tabCRM Meeting` m
		where m.status != 'Cancelled' and m.starts_on >= %(start)s and m.starts_on < %(end)s {cond('m.organizer')} group by 1 order by count desc""",
		params,
		as_dict=True,
	)
	upcoming = []
	for r in frappe.db.sql(
		f"""select m.name, m.subject, m.starts_on, m.reference_docname, m.organizer from `tabCRM Meeting` m
		where m.status = 'Scheduled' and m.starts_on >= now() and m.starts_on < date_add(now(), interval 7 day) {cond('m.organizer')}
		order by m.starts_on asc limit 10""",
		{"who": who},
		as_dict=True,
	):
		title, url = _title(r.reference_docname)
		upcoming.append({"name": r.name, "subject": r.subject, "when": r.starts_on.strftime("%Y-%m-%d %H:%M:%S"), "title": title, "url": url, "by": user_details(r.organizer)[1] or r.organizer})

	outcome_by_person = {}
	for r in frappe.db.sql(
		f"""select f.owner as person, f.outcome as outcome, count(*) as n from `tabCRM Follow Up` f
		where f.mode = 'Call' and f.outcome != '' and f.followed_up_on >= %(start)s and f.followed_up_on < %(end)s {cond('f.owner')}
		group by f.owner, f.outcome""",
		params,
		as_dict=True,
	):
		outcome_by_person.setdefault(r.person or "", {})[r.outcome] = cint(r.n)

	stale_count = 0
	stale_days = cint(get_settings().stale_lead_days)
	if stale_days > 0:
		try:
			stale_count = len([r for r in stale.stale_records("CRM Lead", stale_days) if not who or r.owner_user == who])
		except Exception:
			frappe.log_error(title="CRM Add-ons: could not count stale leads", message=frappe.get_traceback())

	# ---- lists
	tomorrow = add_days(today, 1)
	queue = []
	for r in frappe.db.sql(
		f"""select f.name, f.reference_docname, f.next_follow_up_on, f.assigned_to, f.outcome, f.remark
		from `tabCRM Follow Up` f
		where f.next_closed = 0 and f.next_follow_up_on is not null and f.next_follow_up_on < %(tomorrow)s {cond('f.assigned_to')}
		order by f.next_follow_up_on asc limit 25""",
		{"tomorrow": tomorrow, "who": who},
		as_dict=True,
	):
		title, url = _title(r.reference_docname)
		queue.append(
			{
				"name": r.name,
				"lead": r.reference_docname,
				"title": title,
				"url": url,
				"due": r.next_follow_up_on.strftime("%Y-%m-%d %H:%M:%S"),
				"bucket": "overdue" if r.next_follow_up_on < frappe.utils.get_datetime(today) else "today",
				"owner": user_details(r.assigned_to)[1] or r.assigned_to,
				"outcome": r.outcome,
				"remark": r.remark,
			}
		)

	recent = []
	for r in frappe.db.sql(
		f"""select f.name, f.reference_docname, f.followed_up_on, f.owner, f.outcome, f.remark, f.attempt_no
		from `tabCRM Follow Up` f
		where f.mode = 'Call' and f.followed_up_on >= %(start)s and f.followed_up_on < %(end)s {cond('f.owner')}
		order by f.followed_up_on desc limit 12""",
		params,
		as_dict=True,
	):
		title, url = _title(r.reference_docname)
		recent.append(
			{
				"name": r.name,
				"lead": r.reference_docname,
				"title": title,
				"url": url,
				"when": r.followed_up_on.strftime("%Y-%m-%d %H:%M:%S"),
				"by": user_details(r.owner)[1] or r.owner,
				"outcome": r.outcome,
				"remark": r.remark,
				"attempt": r.attempt_no,
			}
		)

	# ---- people
	team = _team(who)
	names = team | {p for p in current if p} | {p for p in pending if p}
	people = []
	for person in names:
		c = current.get(person) or dict.fromkeys(COUNTS, 0)
		p = pending.get(person) or {}
		tot = _totals(c)
		people.append(
			{
				"user": person,
				"full_name": user_details(person)[1] or person,
				**tot,
				"due_today": p.get("due_today", 0),
				"overdue": p.get("overdue", 0),
				"outcomes": outcome_by_person.get(person, {}),
			}
		)
	people.sort(key=lambda x: (x["calls"], x["follow_ups"]), reverse=True)

	def pending_total(key):
		return sum(v.get(key, 0) for v in pending.values())

	today_totals = _totals(_sum(todays))
	return {
		"scope": {
			"is_manager": manager,
			"user": who,
			"user_name": user_details(who)[1] if who else None,
			"from": str(start),
			"to": str(end),
			"days": span,
			"users": sorted(
				[{"name": u, "full_name": user_details(u)[1] or u} for u in (_team(None) if manager else {frappe.session.user})],
				key=lambda x: x["full_name"].lower(),
			),
		},
		"today": {
			"calls": today_totals["calls"],
			"connected": today_totals["connected"],
			"follow_ups": today_totals["follow_ups"],
			"meetings": today_totals["meetings"],
			"due_today": pending_total("due_today"),
			"overdue": pending_total("overdue"),
		},
		"totals": _totals(_sum(current)),
		"previous": _totals(_sum(previous)),
		"daily": sorted(days.values(), key=lambda x: (x["date"], x["person"])),
		"outcomes": [{"outcome": _(r.outcome), "key": r.outcome, "count": cint(r.count)} for r in outcomes],
		"statuses": [{"status": r.status, "color": r.color, "count": cint(r.count), "age": round(float(r.age or 0))} for r in pipeline],
		"funnel": funnel,
		"speed_hours": speed_hours,
		"by_hour": by_hour,
		"attempts": attempts_rows,
		"sources": [{"source": r.source, "leads": cint(r.leads), "converted": cint(r.converted)} for r in sources],
		"meeting_outcomes": [{"outcome": r.outcome, "count": cint(r.count)} for r in meeting_outcomes],
		"upcoming_meetings": upcoming,
		"stale": stale_count,
		"stale_days": stale_days,
		"score_bands": bands,
		"people": people,
		"queue": queue,
		"recent": recent,
		"now": now_datetime().strftime("%Y-%m-%d %H:%M:%S"),
	}


@frappe.whitelist()
def get_sales_report(from_date=None, to_date=None, user=None):
	frappe.only_for(ROLES)
	return _build_report(from_date, to_date, user)


# -- export ----------------------------------------------------------------------------------------------


def _call_log(start, end, who):
	"""Every call in the period, newest first (up to 10,000), for the detailed sheet / CSV."""
	who_sql = " and f.owner = %(who)s" if who else ""
	rows = frappe.db.sql(
		f"""select f.followed_up_on, f.reference_docname, f.owner, f.outcome, f.attempt_no, f.remark, f.next_follow_up_on
		from `tabCRM Follow Up` f
		where f.mode = 'Call' and f.followed_up_on >= %(start)s and f.followed_up_on < %(end)s {who_sql}
		order by f.followed_up_on desc limit 10000""",
		{"start": start, "end": add_days(end, 1), "who": who},
		as_dict=True,
	)
	titles = {}
	out = []
	for r in rows:
		if r.reference_docname not in titles:
			titles[r.reference_docname] = _title(r.reference_docname)[0]
		out.append(
			[
				r.followed_up_on.strftime("%Y-%m-%d %H:%M"),
				titles[r.reference_docname],
				r.reference_docname,
				user_details(r.owner)[1] or r.owner,
				r.outcome,
				r.attempt_no,
				r.remark or "",
				r.next_follow_up_on.strftime("%Y-%m-%d %H:%M") if r.next_follow_up_on else "",
			]
		)
	return out


def _deals_in_use():
	"""A team that hides the Deals menu is not using Deals: its dashboards and exports leave deal numbers out."""
	return not cint(get_settings().get("hide_deals_menu"))


CALL_LOG_HEADER = ["When", "Lead", "Lead ID", "Called by", "Outcome", "Attempt", "Remark", "Next follow-up"]


def _change(cur, prev):
	if not prev:
		return "new" if cur else ""
	return f"{round((cur - prev) / prev * 100)}%"


def _workbook(data, log):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	head_font = Font(bold=True, color="FFFFFF")
	head_fill = PatternFill("solid", fgColor="1F2937")
	wb = Workbook()

	def sheet(title, header, rows, first=False):
		ws = wb.active if first else wb.create_sheet()
		ws.title = title
		ws.append(header)
		for cell in ws[1]:
			cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(vertical="center")
		for row in rows:
			ws.append(row)
		ws.freeze_panes = "A2"
		for i, col in enumerate(header, 1):
			longest = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r)])
			ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
		return ws

	scope, t, p = data["scope"], data["totals"], data["previous"]
	who = scope["user_name"] if scope["user"] else "Everyone"
	labels = [
		("Calls", "calls"),
		("Connected calls", "connected"),
		("Connect rate (%)", "connect_rate"),
		("Not picked", "not_picked"),
		("Not connected", "not_connected"),
		("Follow-ups added", "follow_ups"),
		("Leads added", "leads"),
		("Meetings", "meetings"),
		("Meetings held", "held"),
		("Converted to deals", "converted"),
		("Attempts to connect (avg)", "avg_attempts"),
	]
	if not _deals_in_use():
		labels = [x for x in labels if x[1] != "converted"]
	summary = [[label, t[k], p[k], _change(t[k], p[k])] for label, k in labels]
	speed = data["speed_hours"] if data["speed_hours"] is not None else ""
	summary += [["Average hours to first call", speed, "", ""], ["Stale leads now", data["stale"], "", ""]]
	ws = sheet("Summary", ["Metric", f"{scope['from']} to {scope['to']}", "Previous period", "Change"], summary, first=True)
	ws.insert_rows(1, 2)
	ws["A1"] = f"Sales Dashboard - {who}"
	ws["A1"].font = Font(bold=True, size=14)
	ws["A2"] = f"Generated {data['now']} by {user_details(frappe.session.user)[1] or frappe.session.user}"
	ws.freeze_panes = "A4"

	people_cols = [
		("Person", "full_name"),
		("Calls", "calls"),
		("Connected", "connected"),
		("Connect rate (%)", "connect_rate"),
		("Not picked", "not_picked"),
		("Not connected", "not_connected"),
		("Follow-ups", "follow_ups"),
		("Leads added", "leads"),
		("Meetings", "meetings"),
		("Meetings held", "held"),
		("Converted", "converted"),
		("Attempts to connect", "avg_attempts"),
		("Due today", "due_today"),
		("Overdue", "overdue"),
	]
	if not _deals_in_use():
		people_cols = [c for c in people_cols if c[1] != "converted"]
	sheet("Team", [c[0] for c in people_cols], [[row[k] for _label, k in people_cols] for row in data["people"]])

	names = {row["user"]: row["full_name"] for row in data["people"]}
	sheet(
		"Daily",
		["Date", "Person", "Calls", "Connected", "Follow-ups", "Leads added"],
		[[r["date"], names.get(r["person"], r["person"]), r["calls"], r["connected"], r["follow_ups"], r["leads"]] for r in data["daily"]],
	)
	total_calls = sum(o["count"] for o in data["outcomes"]) or 1
	sheet(
		"Call outcomes",
		["Outcome", "Calls", "Share (%)"],
		[[o["key"], o["count"], round(o["count"] / total_calls * 100, 1)] for o in data["outcomes"]],
	)
	sheet("Funnel", ["Stage", "Leads"], [[x["stage"], x["count"]] for x in data["funnel"] if _deals_in_use() or x["stage"] != "Converted to deal"])
	sheet("Pipeline", ["Status", "Open leads", "Average age (days)"], [[x["status"], x["count"], x["age"]] for x in data["statuses"]])
	if _deals_in_use():
		sheet("Sources", ["Source", "Leads added", "Converted"], [[x["source"], x["leads"], x["converted"]] for x in data["sources"]])
	else:
		sheet("Sources", ["Source", "Leads added"], [[x["source"], x["leads"]] for x in data["sources"]])
	sheet("Calls by hour", ["Hour", "Calls", "Connected"], [[f"{x['hour']:02d}:00", x["calls"], x["connected"]] for x in data["by_hour"]])
	sheet(
		"Follow-up queue",
		["Lead", "Due", "Bucket", "Owner", "Last outcome", "Remark"],
		[[q["title"], q["due"], q["bucket"], q["owner"], q["outcome"] or "", q["remark"] or ""] for q in data["queue"]],
	)
	sheet("Call log", CALL_LOG_HEADER, log)
	return wb


@frappe.whitelist(methods=["GET"])
def export_report(from_date=None, to_date=None, user=None, format="xlsx"):
	"""Download the dashboard: an Excel workbook (a sheet per section, plus the full call log)
	or the call log as CSV. Same access rules as the dashboard itself."""
	import csv
	import io

	frappe.only_for(ROLES)
	data = _build_report(from_date, to_date, user)
	start, end = getdate(data["scope"]["from"]), getdate(data["scope"]["to"])
	log = _call_log(start, end, data["scope"]["user"])
	stem = f"sales-dashboard_{start}_to_{end}"

	if format == "csv":
		out = io.StringIO()
		writer = csv.writer(out)
		writer.writerow(CALL_LOG_HEADER)
		writer.writerows(log)
		frappe.response["filename"] = f"{stem}_calls.csv"
		frappe.response["filecontent"] = ("﻿" + out.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8
	else:
		out = io.BytesIO()
		_workbook(data, log).save(out)
		frappe.response["filename"] = f"{stem}.xlsx"
		frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "download"


def _campaigns_workbook(data):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")
	wb = Workbook()

	def sheet(title, header, rows, first=False):
		ws = wb.active if first else wb.create_sheet()
		ws.title = title
		ws.append(header)
		for cell in ws[1]:
			cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(vertical="center")
		for row in rows:
			ws.append(row)
		ws.freeze_panes = "A2"
		for i, col in enumerate(header, 1):
			longest = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r)])
			ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
		return ws

	scope, t, p = data["scope"], data["totals"], data["previous"]
	who = "Everyone" if not scope.get("user") else scope["user"]
	labels = [
		("Campaigns", "campaigns"), ("Leads reached", "reached"), ("Messages sent", "sent"), ("Delivered", "delivered"),
		("Read or opened", "read"), ("Replied", "replied"), ("Failed", "failed"), ("Opt-outs", "optouts"),
		("Delivered rate (%)", "delivered_rate"), ("Read rate (%)", "read_rate"), ("Reply rate (%)", "reply_rate"),
	]
	rows = [[label, t.get(k) if t.get(k) is not None else "", p.get(k) if p.get(k) is not None else "", _change(t.get(k) or 0, p.get(k) or 0)] for label, k in labels]
	ws = sheet("Summary", ["Metric", f"{scope['from']} to {scope['to']}", "Previous period", "Change"], rows, first=True)
	ws.insert_rows(1, 2)
	ws["A1"] = f"Campaigns - {who}"
	ws["A1"].font = Font(bold=True, size=14)
	ws["A2"] = f"Generated by {user_details(frappe.session.user)[1] or frappe.session.user}"
	ws.freeze_panes = "A4"
	sheet("By day", ["Date", "Email sent", "WhatsApp sent"], [[d["date"], d["Email"], d["WhatsApp"]] for d in data["by_day"]])
	sheet(
		"Channels", ["Channel", "Sent", "Leads reached", "Delivered", "Read or opened", "Replied", "Failed"],
		[[c, v["sent"], v["reached"], v["delivered"], v["read"], v["replied"], v["failed"]] for c, v in data["channels"].items()],
	)
	sheet(
		"Campaigns", ["Campaign", "Status", "Sent", "Leads reached", "Delivered", "Read or opened", "Replied", "Failed", "Reply rate (%)"],
		[[c["campaign_name"], c["status"], c["sent"], c["reached"], c["delivered"], c["read"], c["replied"], c["failed"], c["reply_rate"] if c["reply_rate"] is not None else ""] for c in data["top_campaigns"]],
	)
	sheet("Status", ["Status", "Campaigns"], [[x["status"], x["count"]] for x in data["status_counts"]])
	return wb


@frappe.whitelist(methods=["GET"])
def export_campaigns(from_date=None, to_date=None, user=None, format="xlsx"):
	"""Download the Campaigns view of the dashboard: an Excel workbook, or the campaign table as CSV.
	Same access rules as the view itself (a sales user only gets their own campaigns)."""
	import csv
	import io

	from crm_addons.dashboard import get_nurturing

	frappe.only_for(ROLES)
	data = get_nurturing(from_date, to_date, user, top_limit=500)
	stem = f"campaigns_{data['scope']['from']}_to_{data['scope']['to']}"
	if format == "csv":
		out = io.StringIO()
		writer = csv.writer(out)
		writer.writerow(["Campaign", "Status", "Sent", "Leads reached", "Delivered", "Read or opened", "Replied", "Failed", "Reply rate (%)"])
		for c in data["top_campaigns"]:
			writer.writerow([c["campaign_name"], c["status"], c["sent"], c["reached"], c["delivered"], c["read"], c["replied"], c["failed"], "" if c["reply_rate"] is None else c["reply_rate"]])
		frappe.response["filename"] = f"{stem}.csv"
		frappe.response["filecontent"] = ("\ufeff" + out.getvalue()).encode("utf-8")
	else:
		out = io.BytesIO()
		_campaigns_workbook(data).save(out)
		frappe.response["filename"] = f"{stem}.xlsx"
		frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "download"
