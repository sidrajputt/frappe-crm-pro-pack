# Copyright (c) 2026, Coding Pro
"""Backend tests for follow-ups (call attempts, outcomes, next follow-up, escalation, queue),
and for the other add-ons built on them: scoring, stale-lead alerts, meeting outcomes,
templates and WhatsApp.

Each test runs in a transaction that is rolled back.
"""

import json
from unittest.mock import patch

import frappe
from frappe.utils import add_to_date, now_datetime

from crm_addons import api, dashboard, followups, install, reports, scoring, stale
from crm_addons.meetings.doctype.crm_meeting.test_crm_meeting import FakeGoogle, google_patch
from crm_addons.meetings.doctype.crm_meeting.test_crm_meeting import TestCase as MeetingsTestCase
from crm_addons.patches import follow_up_outcomes


class FollowUpBase(MeetingsTestCase):
	def setUp(self):
		super().setUp()
		frappe.db.set_value(
			"CRM Addons Settings",
			None,
			{
				"enable_follow_ups": 1,
				"follow_up_reminders": 1,
				"escalate_after_attempts": 3,
				"unreachable_lead_status": "",
				"lead_scoring_enabled": 1,
				"stale_lead_days": 7,
			},
		)
		frappe.clear_document_cache("CRM Addons Settings", "CRM Addons Settings")
		self.clock = 0

	def log(self, outcome="Did Not Pick", **kw):
		"""Log a call one hour after the previous one, so the order is unambiguous."""
		self.clock += 1
		data = {
			"reference_doctype": "CRM Lead",
			"reference_docname": self.lead.name,
			"outcome": outcome,
			"followed_up_on": f"2020-01-01 {9 + self.clock:02d}:00:00",
			**kw,
		}
		return followups.save_follow_up(frappe.as_json(data))

	def rows(self):
		return frappe.get_all(
			"CRM Follow Up",
			filters={"reference_docname": self.lead.name},
			fields=["name", "attempt_no", "call_status", "outcome", "next_closed", "next_follow_up_on", "reminded"],
			order_by="followed_up_on asc",
		)

	def lead_value(self, fieldname):
		return frappe.db.get_value("CRM Lead", self.lead.name, fieldname)


class TestFollowUps(FollowUpBase):
	def test_attempts_count_calls_in_a_row_and_restart_after_a_connect(self):
		for outcome in ("Did Not Pick", "Did Not Connect", "Interested", "Did Not Pick"):
			self.log(outcome)
		rows = self.rows()
		self.assertEqual([r.attempt_no for r in rows], [1, 2, 3, 1])
		self.assertEqual([r.call_status for r in rows], ["Did Not Pick", "Did Not Connect", "Connected", "Did Not Pick"])

	def test_a_call_needs_an_outcome(self):
		with self.assertRaises(frappe.ValidationError):
			self.log("")

	def test_only_leads_can_have_follow_ups(self):
		with self.assertRaises(frappe.ValidationError):
			self.log(reference_doctype="CRM Deal", reference_docname="anything")

	def test_newest_follow_up_closes_the_older_one(self):
		self.log("Did Not Pick", next_follow_up_on="2020-01-05 10:00:00")
		self.log("Call Back Later", next_follow_up_on="2020-01-09 10:00:00")
		first, second = self.rows()
		self.assertEqual((first.next_closed, second.next_closed), (1, 0))
		self.assertEqual(str(self.lead_value("next_follow_up_on")), "2020-01-09 10:00:00")

	def test_no_next_date_means_nothing_is_open(self):
		self.log("Not Interested")
		self.assertEqual(self.rows()[0].next_closed, 1)
		self.assertIsNone(self.lead_value("next_follow_up_on"))

	def test_next_follow_up_cannot_be_before_the_call(self):
		with self.assertRaises(frappe.ValidationError):
			self.log("Did Not Pick", next_follow_up_on="2019-12-31 10:00:00")

	def test_lead_columns_show_the_latest_outcome_remark_and_attempt(self):
		self.log("Did Not Pick", remark="Phone was ringing")
		self.log("Interested", remark="Wants a quote", next_follow_up_on="2020-01-04 10:00:00")
		self.assertEqual(self.lead_value("addons_last_outcome"), "Interested")
		self.assertEqual(self.lead_value("addons_last_remark"), "Wants a quote")
		self.assertEqual(self.lead_value("addons_call_attempts"), 2)  # connected on the second attempt

	def test_timeline_comment_and_summary(self):
		self.log("Did Not Pick", remark="Phone was ringing")
		result = self.log("Did Not Connect")
		self.assertTrue(any("Phone was ringing" in c for c in self.comments()))
		self.assertEqual(result["summary"]["streak"], 2)
		self.assertEqual(result["summary"]["last_status"], "Did Not Connect")

	def test_the_form_can_change_the_lead_status(self):
		statuses = frappe.get_all("CRM Lead Status", pluck="name", order_by="position asc")
		current = self.lead_value("status")
		other = next((s for s in statuses if s != current), None)
		if not other:
			self.skipTest("the site has only one lead status")
		self.log("Interested", lead_status=other)
		self.assertEqual(self.lead_value("status"), other)
		with self.assertRaises(frappe.ValidationError):
			self.log("Interested", lead_status="No Such Status")

	def test_form_data_has_what_the_form_needs(self):
		self.log("Did Not Pick")
		self.log("Did Not Connect")
		data = followups.get_form_data("CRM Lead", self.lead.name)
		self.assertEqual(data["streak"], 2)  # the next call is attempt 3
		self.assertEqual(data["record"]["status"], self.lead_value("status"))
		self.assertEqual(
			data["config"]["outcomes"],
			["Did Not Pick", "Did Not Connect", "Interested", "Not Interested", "Meeting Scheduled", "Ask for Detail", "Call Back Later"],
		)
		self.assertTrue(data["statuses"])
		entry = followups.get_form_data("CRM Lead", self.lead.name, name=self.rows()[0].name)["entry"]
		self.assertEqual(entry["outcome"], "Did Not Pick")

	def test_escalates_once_after_the_configured_number_of_missed_calls(self):
		before = frappe.db.count("CRM Notification")
		self.log("Did Not Pick")
		self.log("Did Not Connect")
		self.assertFalse(any("No answer after" in c for c in self.comments()))
		self.log("Did Not Pick")
		self.assertTrue(any("No answer after 3 calls" in c for c in self.comments()))
		self.assertGreater(frappe.db.count("CRM Notification"), before)
		self.log("Did Not Pick")  # a fourth call does not escalate again
		self.assertEqual(sum("No answer after" in c for c in self.comments()), 1)

	def test_escalation_can_be_switched_off(self):
		frappe.db.set_value("CRM Addons Settings", None, "escalate_after_attempts", 0)
		frappe.clear_document_cache("CRM Addons Settings", "CRM Addons Settings")
		for _i in range(3):
			self.log("Did Not Pick")
		self.assertFalse(any("No answer after" in c for c in self.comments()))

	def test_queue_buckets(self):
		now = now_datetime()
		self.log("Did Not Pick", followed_up_on="2020-01-01 09:00:00", next_follow_up_on="2020-01-02 09:00:00")
		self.assertEqual(followups.get_queue(scope="all")["counts"], {"overdue": 1, "today": 0, "upcoming": 0})

		self.log("Did Not Pick", followed_up_on=str(add_to_date(now, minutes=-5)), next_follow_up_on=str(add_to_date(now, days=5)))
		queue = followups.get_queue(scope="all")
		self.assertEqual(queue["counts"], {"overdue": 0, "today": 0, "upcoming": 1})  # the older one was closed
		self.assertEqual(queue["items"][0]["bucket"], "upcoming")

	def test_complete_removes_it_from_the_queue(self):
		self.log("Did Not Pick", next_follow_up_on="2020-01-02 09:00:00")
		followups.complete_follow_up(self.rows()[0].name)
		self.assertEqual(followups.get_queue(scope="all")["counts"]["overdue"], 0)
		self.assertIsNone(self.lead_value("next_follow_up_on"))

	def test_reminder_is_sent_once_when_due(self):
		self.log("Did Not Pick", next_follow_up_on="2020-01-02 09:00:00")
		before = frappe.db.count("CRM Notification")
		with patch.object(frappe, "sendmail"), patch.object(frappe.db, "commit"):  # the scheduler commits
			followups.send_due_reminders()
			after_first = frappe.db.count("CRM Notification")
			followups.send_due_reminders()
		self.assertEqual(after_first, before + 1)
		self.assertEqual(frappe.db.count("CRM Notification"), after_first)
		self.assertEqual(self.rows()[0].reminded, 1)

	def test_reminder_survives_an_email_that_cannot_be_sent(self):
		frappe.db.set_value("CRM Addons Settings", None, "email_reminders", 1)
		frappe.clear_document_cache("CRM Addons Settings", "CRM Addons Settings")
		self.log("Did Not Pick", next_follow_up_on="2020-01-02 09:00:00")
		before = frappe.db.count("CRM Notification")
		# rollback is patched too, or it would undo the whole test transaction
		with patch.object(frappe, "sendmail", side_effect=Exception("no outgoing email account")), patch.object(
			frappe.db, "commit"
		), patch.object(frappe.db, "rollback"):
			followups.send_due_reminders()
		self.assertEqual(frappe.db.count("CRM Notification"), before + 1)
		self.assertEqual(self.rows()[0].reminded, 1)
		self.assertIn("email failed", self.log_error.call_args.kwargs["title"])

	def test_sales_users_only_see_their_own_follow_ups(self):
		self.log("Did Not Pick")
		other = self.make_user("other.sales@example.org", "Sales User")
		frappe.set_user(other.name)
		self.assertEqual(frappe.get_list("CRM Follow Up", fields=["name"]), [])
		frappe.set_user("Administrator")
		self.assertEqual(len(frappe.get_list("CRM Follow Up", fields=["name"])), 1)

	def test_dashboard_charts_return_the_shapes_crm_expects(self):
		self.log("Did Not Pick", followed_up_on=str(now_datetime()))
		self.log("Interested", followed_up_on=str(add_to_date(now_datetime(), minutes=1)))
		start, end = "2000-01-01", "2100-01-01"
		for fn in (
			dashboard.get_follow_ups_due_today,
			dashboard.get_overdue_follow_ups,
			dashboard.get_call_connect_rate,
			dashboard.get_attempts_to_connect,
		):
			self.assertIn("value", fn(start, end, None))
		self.assertEqual(dashboard.get_call_connect_rate(start, end, None)["value"], 50.0)
		self.assertEqual(dashboard.get_attempts_to_connect(start, end, None)["value"], 2.0)
		self.assertEqual(dashboard.get_follow_ups_by_day(start, end, "Administrator")["xAxis"]["type"], "time")
		self.assertIn("categoryColumn", dashboard.get_follow_ups_by_outcome(start, end, None))
		self.assertIn("categoryColumn", dashboard.get_calls_by_status(start, end, None))


class TestAddonsExtras(FollowUpBase):
	"""Scoring, stale-lead alerts, meeting outcomes, templates, WhatsApp and the data patch."""

	def test_score_bands(self):
		self.assertEqual([scoring.band(x) for x in (0, 29, 30, 59, 60, 100)], ["Cold", "Cold", "Warm", "Warm", "Hot", "Hot"])

	def test_score_follows_the_calls_and_the_latest_outcome(self):
		self.log("Interested")
		self.assertEqual(self.lead_value("addons_score"), 20)
		self.log("Not Interested")
		self.assertEqual(self.lead_value("addons_score"), 0)  # never below 0
		self.assertEqual(followups.list_follow_ups("CRM Lead", self.lead.name)["score_band"], "Cold")

	def test_stale_lead_is_alerted_once(self):
		status = frappe.db.get_value("CRM Lead Status", self.lead_value("status"), "type")
		if status not in ("Open", "Ongoing", "OnHold"):
			self.skipTest("the test lead is not in an open status")
		frappe.db.set_value("CRM Lead", self.lead.name, "modified", "2020-01-01 00:00:00", update_modified=False)
		# creating the lead also posts timeline comments, which count as activity
		frappe.db.set_value(
			"Comment",
			{"reference_doctype": "CRM Lead", "reference_name": self.lead.name},
			"creation",
			"2020-01-01 00:00:00",
			update_modified=False,
		)
		self.assertIn(self.lead.name, [r.name for r in stale.stale_records("CRM Lead", 7)])

		before = frappe.db.count("CRM Notification")
		with patch.object(frappe.db, "commit"):  # the scheduler commits
			stale.send_stale_alerts()
			after_first = frappe.db.count("CRM Notification")
			stale.send_stale_alerts()
		self.assertGreater(after_first, before)
		self.assertEqual(frappe.db.count("CRM Notification"), after_first)

	def test_meeting_outcome_closes_the_meeting_and_logs_a_follow_up(self):
		with patch.object(frappe, "sendmail"):
			meeting = self.save(provider="Manual Link", add_video_conferencing=0)["meeting"]
		result = api.record_meeting_outcome(meeting["name"], "Held", "Sent the quote", "2030-01-05 10:00:00")
		self.assertEqual((result["status"], result["meeting_outcome"]), ("Completed", "Held"))
		self.assertTrue(any("Meeting held" in c for c in self.comments()))
		entry = self.rows()[0]
		self.assertEqual(str(entry.next_follow_up_on), "2030-01-05 10:00:00")
		with self.assertRaises(frappe.ValidationError):
			api.record_meeting_outcome(meeting["name"], "Maybe")

	def test_rescheduled_outcome_keeps_the_meeting_open(self):
		with patch.object(frappe, "sendmail"):
			meeting = self.save(provider="Manual Link", add_video_conferencing=0)["meeting"]
		self.assertEqual(api.record_meeting_outcome(meeting["name"], "Rescheduled")["status"], "Scheduled")

	def test_desk_workspace_is_created_once_and_left_to_the_admin(self):
		if frappe.db.exists("Workspace", install.WORKSPACE):
			frappe.delete_doc("Workspace", install.WORKSPACE, force=1, ignore_permissions=True)
		install.ensure_workspace()
		if not frappe.db.exists("Workspace", install.WORKSPACE):
			self.skipTest("this Frappe version does not let the app create a workspace (it is optional)")
		workspace = frappe.get_doc("Workspace", install.WORKSPACE)
		labels = [label for label, _kind, _target in install.WORKSPACE_SHORTCUTS]
		self.assertEqual([row.label for row in workspace.shortcuts], labels)
		blocks = json.loads(workspace.content)
		self.assertEqual(blocks[0]["type"], "header")
		self.assertEqual({b["data"]["shortcut_name"] for b in blocks if b["type"] == "shortcut"}, set(labels))
		workspace.db_set("title", "My own title")
		install.ensure_workspace()  # the admin renamed it: it is not rewritten
		self.assertEqual(frappe.db.get_value("Workspace", install.WORKSPACE, "title"), "My own title")

	def test_starter_templates_are_added_once(self):
		frappe.db.set_default("crm_addons_templates_seeded", "")
		install.ensure_templates()
		self.assertTrue(frappe.db.exists("CRM Follow Up Template", "Did not pick - try tomorrow"))
		frappe.delete_doc("CRM Follow Up Template", "Did not pick - try tomorrow", force=1)
		install.ensure_templates()  # the admin deleted it: it does not come back
		self.assertFalse(frappe.db.exists("CRM Follow Up Template", "Did not pick - try tomorrow"))
		templates = {t["title"]: t for t in followups.get_config()["templates"]}
		self.assertEqual(templates["Call back later"]["outcome"], "Call Back Later")

	def test_whatsapp_goes_through_crms_own_integration(self):
		frappe.db.set_value("CRM Lead", self.lead.name, "mobile_no", "+911234567890")
		with patch.object(followups, "_whatsapp_ready", return_value=True), patch(
			"crm.api.whatsapp.create_whatsapp_message", return_value="WA-1"
		) as send:
			result = followups.send_missed_call_whatsapp("CRM Lead", self.lead.name, "Sorry I missed you")
		self.assertEqual(result, {"message": "WA-1", "to": "+911234567890"})
		self.assertEqual(send.call_args.args, ("CRM Lead", self.lead.name, "Sorry I missed you", "+911234567890", "", ""))

	def test_whatsapp_needs_a_number_and_a_message(self):
		with patch.object(followups, "_whatsapp_ready", return_value=True):
			with self.assertRaises(frappe.ValidationError):
				followups.send_missed_call_whatsapp("CRM Lead", self.lead.name, "Hello")
			frappe.db.set_value("CRM Lead", self.lead.name, "mobile_no", "+911234567890")
			with self.assertRaises(frappe.ValidationError):
				followups.send_missed_call_whatsapp("CRM Lead", self.lead.name, "  ")

	def test_new_charts_return_the_shapes_crm_expects(self):
		self.log("Interested", followed_up_on=str(now_datetime()))
		start, end = "2000-01-01", "2100-01-01"
		self.assertEqual(dashboard.get_activity_by_salesperson(start, end, None)["xAxis"]["type"], "category")
		self.assertIn("categoryColumn", dashboard.get_leads_by_score_band(start, end, None))
		self.assertIn("categoryColumn", dashboard.get_meetings_by_outcome(start, end, None))

	def test_patch_converts_the_old_status_and_outcome_values(self):
		self.log("Interested")
		self.log("Interested")
		first, second = [r.name for r in self.rows()]
		# how an earlier version stored them
		frappe.db.set_value("CRM Follow Up", first, {"call_status": "Ringing", "outcome": ""})
		frappe.db.set_value("CRM Follow Up", second, {"call_status": "Connected", "outcome": "Send Details"})
		with patch.object(frappe.db, "commit"):
			follow_up_outcomes.execute()
		rows = {r.name: r for r in self.rows()}
		self.assertEqual((rows[first].outcome, rows[first].call_status, rows[first].attempt_no), ("Did Not Pick", "Did Not Pick", 1))
		self.assertEqual((rows[second].outcome, rows[second].call_status, rows[second].attempt_no), ("Ask for Detail", "Connected", 2))


class TestSalesReport(FollowUpBase):
	"""The Sales Dashboard: numbers, and who is allowed to see whose."""

	def setUp(self):
		super().setUp()
		self.today = str(now_datetime().date())
		self.log("Did Not Pick", followed_up_on=str(now_datetime()))
		self.log("Interested", followed_up_on=str(add_to_date(now_datetime(), minutes=1)))
		# the second call was made by another salesperson
		self.rep = self.owner
		frappe.db.set_value("CRM Follow Up", self.rows()[1].name, "owner", self.rep.name)

	def report(self, **kw):
		return reports.get_sales_report(self.today, self.today, **kw)

	def test_manager_sees_everyone(self):
		data = self.report()
		self.assertTrue(data["scope"]["is_manager"])
		self.assertEqual((data["totals"]["calls"], data["totals"]["connected"], data["totals"]["connect_rate"]), (2, 1, 50.0))
		self.assertEqual(data["totals"]["follow_ups"], 2)
		self.assertEqual(sum(r["calls"] for r in data["daily"]), 2)
		self.assertEqual({o["key"]: o["count"] for o in data["outcomes"]}, {"Did Not Pick": 1, "Interested": 1})
		by_user = {p["user"]: p for p in data["people"]}
		self.assertEqual(by_user["Administrator"]["calls"], 1)
		self.assertEqual(by_user[self.rep.name]["connected"], 1)

	def test_manager_can_look_at_one_person(self):
		data = self.report(user=self.rep.name)
		self.assertEqual(data["totals"]["calls"], 1)
		self.assertEqual([p["user"] for p in data["people"]], [self.rep.name])

	def test_sales_user_only_ever_sees_their_own_numbers(self):
		frappe.set_user(self.rep.name)
		for asked in (None, "Administrator", self.rep.name):
			data = self.report(user=asked)
			self.assertFalse(data["scope"]["is_manager"])
			self.assertEqual(data["scope"]["user"], self.rep.name)
			self.assertEqual(data["totals"]["calls"], 1)
			self.assertEqual([p["user"] for p in data["people"]], [self.rep.name])
			self.assertEqual([u["name"] for u in data["scope"]["users"]], [self.rep.name])

	def test_users_without_a_sales_role_cannot_open_it(self):
		self.addCleanup(setattr, frappe.flags, "in_test", frappe.flags.in_test)
		frappe.flags.in_test = False  # Frappe's only_for() lets everyone through while tests run; check it for real
		nobody = self.make_user("nobody.test@example.org", "Blogger")
		frappe.set_user(nobody.name)
		with self.assertRaises(frappe.PermissionError):
			self.report()

	def test_queue_and_today_counts(self):
		follow_up = self.rows()[0].name
		frappe.db.set_value("CRM Follow Up", follow_up, {"next_closed": 0, "next_follow_up_on": str(add_to_date(now_datetime(), minutes=2))})
		data = self.report()
		self.assertEqual(data["today"]["due_today"], 1)
		self.assertEqual([q["bucket"] for q in data["queue"]], ["today"])
		self.assertEqual(data["queue"][0]["title"], "Meeting Tester")

	def test_period_is_clamped_and_reversed_dates_are_fixed(self):
		start, end = reports._period("2030-01-10", "2030-01-01")
		self.assertEqual((str(start), str(end)), ("2030-01-01", "2030-01-10"))
		start, end = reports._period("2000-01-01", "2030-01-01")
		self.assertEqual((end - start).days, 366)

	def test_report_has_every_section_the_dashboard_draws(self):
		data = self.report()
		for key in (
			"scope", "today", "totals", "previous", "daily", "outcomes", "statuses", "funnel", "speed_hours", "by_hour",
			"attempts", "sources", "meeting_outcomes", "upcoming_meetings", "stale", "stale_days", "score_bands", "people", "queue", "recent",
		):
			self.assertIn(key, data)
		self.assertEqual(len(data["by_hour"]), 24)
		self.assertEqual([a["label"] for a in data["attempts"]], ["1st call", "2nd call", "3rd call", "4th or later"])
		self.assertEqual(data["attempts"][1]["count"], 1)  # the connected call was the second attempt
		self.assertEqual([s["stage"] for s in data["funnel"]][0], "Leads added")
		self.assertEqual(
			{p["user"]: p["outcomes"] for p in data["people"]}["Administrator"], {"Did Not Pick": 1}
		)

	def test_funnel_counts_what_became_of_new_leads(self):
		data = self.report()
		stages = {s["stage"]: s["count"] for s in data["funnel"]}
		self.assertGreaterEqual(stages["Leads added"], 1)  # the test lead was created today
		self.assertGreaterEqual(stages["Contacted"], 1)
		self.assertGreaterEqual(stages["Reached on a call"], 1)
		self.assertLessEqual(stages["Reached on a call"], stages["Contacted"])

	def test_excel_export_has_a_sheet_per_section_and_the_call_log(self):
		import io

		from openpyxl import load_workbook

		frappe.response.clear()
		reports.export_report(self.today, self.today)
		self.assertEqual(frappe.response["type"], "download")
		self.assertTrue(frappe.response["filename"].endswith(".xlsx"))
		wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
		self.assertEqual(
			wb.sheetnames,
			["Summary", "Team", "Daily", "Call outcomes", "Funnel", "Pipeline", "Sources", "Calls by hour", "Follow-up queue", "Call log"],
		)
		log = list(wb["Call log"].iter_rows(values_only=True))
		self.assertEqual(log[0][:5], ("When", "Lead", "Lead ID", "Called by", "Outcome"))
		self.assertEqual(len(log) - 1, 2)  # both calls
		self.assertEqual(wb["Summary"]["A1"].value.startswith("Sales Dashboard"), True)

	def test_csv_export_is_the_call_log(self):
		frappe.response.clear()
		reports.export_report(self.today, self.today, format="csv")
		text = frappe.response["filecontent"].decode("utf-8-sig")
		self.assertTrue(frappe.response["filename"].endswith("_calls.csv"))
		lines = text.strip().splitlines()
		self.assertEqual(lines[0], ",".join(reports.CALL_LOG_HEADER))
		self.assertEqual(len(lines), 3)

	def test_sales_user_export_only_contains_their_own_calls(self):
		frappe.set_user(self.rep.name)
		frappe.response.clear()
		reports.export_report(self.today, self.today, user="Administrator", format="csv")
		lines = frappe.response["filecontent"].decode("utf-8-sig").strip().splitlines()
		self.assertEqual(len(lines), 2)  # header and the one call this person made


class TestFollowUpWorkspace(FollowUpBase):
	"""The Follow-ups page: filters, tabs, numbers, reschedule and bulk actions, and who sees whose."""

	def setUp(self):
		super().setUp()
		now = now_datetime()
		# one lead, one open follow-up (the newest closes the earlier ones)
		self.log("Did Not Pick", followed_up_on=str(add_to_date(now, days=-5)), next_follow_up_on=str(add_to_date(now, days=-1)))
		self.rep = self.owner
		self.other_lead = frappe.get_doc(
			{"doctype": "CRM Lead", "first_name": "Zed", "last_name": "Other", "email": "zed.other@example.com", "mobile_no": "5550101", "lead_owner": "Administrator"}
		).insert(ignore_permissions=True)
		frappe.db.set_value("CRM Follow Up", self.rows()[0].name, "assigned_to", "Administrator")
		self.mine = followups.save_follow_up(
			frappe.as_json({
				"reference_doctype": "CRM Lead", "reference_docname": self.other_lead.name, "outcome": "Interested",
				"followed_up_on": str(add_to_date(now, days=-2)), "next_follow_up_on": str(add_to_date(now, days=3)),
				"assigned_to": self.rep.name, "remark": "Wants a quote",
			})
		)

	def workspace(self, **kw):
		return followups.get_workspace(**kw)

	def test_manager_sees_everyone_by_default(self):
		data = self.workspace()
		self.assertEqual(data["scope"]["value"], "all")
		self.assertEqual(data["counts"]["overdue"], 1)
		self.assertEqual(data["counts"]["upcoming"], 1)
		self.assertEqual({o["name"] for o in data["owners"]}, {"Administrator", self.rep.name})
		self.assertEqual(data["items"][0]["bucket"], "overdue")
		self.assertEqual(data["items"][0]["reference_title"], "Meeting Tester")

	def test_manager_can_narrow_to_one_owner_or_to_themselves(self):
		self.assertEqual(self.workspace(tab="upcoming", owner=self.rep.name)["total"], 1)
		self.assertEqual(self.workspace(tab="upcoming", owner="Administrator")["total"], 0)
		self.assertEqual(self.workspace(scope="mine", tab="overdue")["total"], 1)
		self.assertEqual(self.workspace(scope="mine", tab="upcoming")["total"], 0)

	def test_sales_user_only_gets_their_own_whatever_is_asked(self):
		frappe.set_user(self.rep.name)
		for kw in ({}, {"scope": "all"}, {"owner": "Administrator"}, {"scope": "all", "owner": "Administrator"}):
			data = self.workspace(tab="upcoming", **kw)
			self.assertEqual(data["scope"]["value"], "mine")
			self.assertFalse(data["scope"]["is_manager"])
			self.assertEqual([str(i["name"]) for i in data["items"]], [str(self.mine["items"][0]["name"])])
			self.assertEqual(data["owners"], [])
			self.assertEqual(data["counts"]["overdue"], 0)

	def test_search_matches_lead_name_phone_and_remark(self):
		self.assertEqual(self.workspace(tab="upcoming", search="Zed")["total"], 1)
		self.assertEqual(self.workspace(tab="upcoming", search="5550101")["total"], 1)
		self.assertEqual(self.workspace(tab="upcoming", search="quote")["total"], 1)
		self.assertEqual(self.workspace(tab="upcoming", search="nobody-has-this")["total"], 0)
		self.assertEqual(self.workspace(tab="overdue", search="Zed")["total"], 0)

	def test_outcome_filter_and_row_details(self):
		data = self.workspace(tab="upcoming", outcome="Interested")
		self.assertEqual(data["total"], 1)
		item = data["items"][0]
		self.assertEqual((item["phone"], item["outcome"], item["remark"]), ("5550101", "Interested", "Wants a quote"))
		self.assertEqual(self.workspace(tab="upcoming", outcome="Did Not Pick")["total"], 0)

	def test_date_range_limits_the_due_date(self):
		today = str(now_datetime().date())
		self.assertEqual(self.workspace(tab="overdue", from_date=today, to_date=today)["total"], 0)
		wide = self.workspace(tab="upcoming", from_date=str(add_to_date(now_datetime(), days=-10).date()), to_date=str(add_to_date(now_datetime(), days=10).date()))
		self.assertEqual(wide["total"], 1)
		self.assertFalse(wide["done_window"]["default"])

	def test_done_tab_lists_closed_follow_ups(self):
		followups.complete_follow_up(self.mine["items"][0]["name"])
		data = self.workspace(tab="done", from_date=str(add_to_date(now_datetime(), days=1).date()), to_date=str(add_to_date(now_datetime(), days=5).date()))
		self.assertEqual(data["total"], 1)
		self.assertEqual(data["items"][0]["bucket"], "done")
		self.assertEqual(self.workspace()["counts"]["upcoming"], 0)

	def test_paging(self):
		first = self.workspace(tab="overdue", page=1, page_length=1)
		self.assertEqual((len(first["items"]), first["total"]), (1, 1))
		self.assertEqual(self.workspace(tab="overdue", page=2, page_length=1)["items"], [])
		self.assertEqual(self.workspace(page_length=100000)["page_length"], followups.PAGE_MAX)

	def test_reschedule_moves_the_date_and_reopens_the_reminder(self):
		name = self.mine["items"][0]["name"]
		when = str(add_to_date(now_datetime(), days=9))
		followups.reschedule_follow_up(name, when)
		doc = frappe.get_doc("CRM Follow Up", name)
		self.assertEqual(str(doc.next_follow_up_on)[:10], when[:10])
		self.assertEqual((doc.next_closed, doc.reminded), (0, 0))
		self.assertTrue(any("rescheduled" in c for c in self.comments_of(self.other_lead.name)))

	def test_reschedule_rejects_done_and_past_dates(self):
		name = self.mine["items"][0]["name"]
		with self.assertRaises(frappe.ValidationError):
			followups.reschedule_follow_up(name, "2000-01-01 10:00:00")
		followups.complete_follow_up(name)
		with self.assertRaises(frappe.ValidationError):
			followups.reschedule_follow_up(name, str(add_to_date(now_datetime(), days=9)))

	def test_a_sales_user_cannot_reschedule_somebody_elses_follow_up(self):
		stranger = self.make_user("stranger.test@example.org", "Sales User")
		frappe.set_user(stranger.name)
		with self.assertRaises(frappe.PermissionError):
			followups.reschedule_follow_up(self.mine["items"][0]["name"], str(add_to_date(now_datetime(), days=9)))

	def test_bulk_done_and_reschedule_report_what_failed(self):
		a, b = str(self.rows()[0].name), str(self.mine["items"][0]["name"])
		res = followups.bulk_update(frappe.as_json([a, b, "999999"]), "done")
		self.assertEqual(sorted(res["updated"]), sorted([a, b]))
		self.assertEqual([f["name"] for f in res["failed"]], ["999999"])
		self.assertEqual(self.workspace()["counts"]["overdue"] + self.workspace()["counts"]["upcoming"], 0)
		res = followups.bulk_update([a], "reschedule", str(add_to_date(now_datetime(), days=4)))
		self.assertEqual(res["updated"], [])  # already done
		self.assertEqual(len(res["failed"]), 1)
		with self.assertRaises(frappe.ValidationError):
			followups.bulk_update([a], "delete")

	def test_bulk_is_checked_per_follow_up_for_sales_users(self):
		frappe.set_user(self.rep.name)
		res = followups.bulk_update([str(self.rows()[0].name), str(self.mine["items"][0]["name"])], "done")
		self.assertEqual(res["updated"], [str(self.mine["items"][0]["name"])])
		self.assertEqual(len(res["failed"]), 1)

	def test_get_queue_is_unchanged(self):
		data = followups.get_queue("all")
		self.assertEqual(data["counts"], {"overdue": 1, "today": 0, "upcoming": 1})

	def comments_of(self, lead):
		return [c.content for c in frappe.get_all("Comment", filters={"reference_doctype": "CRM Lead", "reference_name": lead}, fields=["content"])]


class TestNurturing(FollowUpBase):
	"""Lead Nurturing: campaign numbers for the Sales Dashboard, and whose campaigns each person sees."""

	def setUp(self):
		super().setUp()
		# the numbers below assume email opens are not tracked; a site may have an outgoing account that tracks them
		frappe.db.set_value("Email Account", {"track_email_status": 1}, "track_email_status", 0)
		self.rep = self.owner
		self.rival = self.make_user("rival.test@example.org", "Sales User")
		self.today = str(now_datetime().date())
		self.mine = self.campaign("Mine Camp", self.rep.name)
		self.theirs = self.campaign("Theirs Camp", self.rival.name)
		sent = now_datetime()
		self.recipient(self.mine, "Email", "Sent", sent)
		self.recipient(self.mine, "WhatsApp", "Read", sent, lead=self.lead)
		self.recipient(self.mine, "WhatsApp", "Failed", None, failed_at=sent)
		self.recipient(self.theirs, "WhatsApp", "Delivered", sent)
		self.recipient(self.theirs, "Email", "Sent", sent)

	def campaign(self, name, owner):
		doc = frappe.get_doc(
			{"doctype": "CRM Campaign", "campaign_name": name, "campaign_owner": owner, "status": "Completed", "audience_mode": "Selected Records", "selected_records": "[]"}
		)
		doc.flags.ignore_validate = True
		doc.flags.ignore_mandatory = True
		doc.insert(ignore_permissions=True)
		frappe.db.set_value("CRM Campaign", doc.name, "owner", owner)
		return doc.name

	def recipient(self, campaign, channel, status, sent_at, lead=None, **kw):
		lead = lead or frappe.get_doc(
			{"doctype": "CRM Lead", "first_name": f"N{frappe.generate_hash(length=6)}", "last_name": "Nurture", "lead_owner": self.rep.name}
		).insert(ignore_permissions=True)
		doc = frappe.get_doc(
			{
				"doctype": "CRM Campaign Recipient", "campaign": campaign, "channel": channel, "status": status, "recipient_type": "CRM Lead",
				"recipient_id": lead.name, "recipient_name": lead.lead_name, "sent_at": sent_at, "step_idx": 1,
				"dedupe_key": frappe.generate_hash(length=20), **kw,
			}
		)
		doc.flags.ignore_validate = True
		doc.flags.ignore_mandatory = True
		doc.insert(ignore_permissions=True)
		return doc

	def nurturing(self, **kw):
		return dashboard.get_nurturing(self.today, self.today, **kw)

	def test_manager_sees_every_campaign(self):
		data = self.nurturing()
		t = data["totals"]
		self.assertEqual((t["campaigns"], t["sent"], t["failed"], t["reached"]), (2, 4, 1, 4))
		self.assertEqual(data["channels"]["WhatsApp"]["sent"], 2)
		self.assertEqual(data["channels"]["WhatsApp"]["delivered"], 2)  # Delivered and Read both count
		self.assertEqual(data["channels"]["WhatsApp"]["read"], 1)
		self.assertEqual(t["delivered_rate"], 100.0)
		self.assertEqual(t["read_rate"], 50.0)  # email opens are not tracked on this site
		self.assertEqual(sum(d["Email"] + d["WhatsApp"] for d in data["by_day"]), 4)
		self.assertEqual({r["status"]: r["count"] for r in data["status_counts"]}, {"Completed": 2})
		self.assertEqual({c["campaign_name"] for c in data["top_campaigns"]}, {"Mine Camp", "Theirs Camp"})
		self.assertTrue(data["has_campaigns"])

	def test_manager_can_look_at_one_person(self):
		data = self.nurturing(user=self.rep.name)
		self.assertEqual((data["totals"]["campaigns"], data["totals"]["sent"]), (1, 2))
		self.assertEqual([c["campaign_name"] for c in data["top_campaigns"]], ["Mine Camp"])

	def test_sales_user_only_sees_their_own_campaigns(self):
		frappe.set_user(self.rep.name)
		for asked in (None, "Administrator", self.rival.name, self.rep.name):
			data = self.nurturing(user=asked)
			self.assertFalse(data["scope"]["is_manager"])
			self.assertEqual(data["scope"]["user"], self.rep.name)
			self.assertEqual((data["totals"]["campaigns"], data["totals"]["sent"]), (1, 2))
			self.assertEqual([c["campaign_name"] for c in data["top_campaigns"]], ["Mine Camp"])
			self.assertNotIn("Theirs Camp", str(data))

	def test_replies_and_opt_outs_are_counted(self):
		frappe.get_doc(
			{"doctype": "Communication", "communication_type": "Communication", "sent_or_received": "Received", "reference_doctype": "CRM Lead",
			 "reference_name": self.lead.name, "subject": "Re: hi", "content": "Interested", "sender": "x@example.com", "creation": add_to_date(now_datetime(), minutes=5)}
		).insert(ignore_permissions=True)
		frappe.db.set_value("Communication", {"reference_name": self.lead.name}, "creation", add_to_date(now_datetime(), minutes=5))
		data = self.nurturing()
		self.assertEqual(data["totals"]["replied"], 1)
		self.assertEqual(data["totals"]["reply_rate"], 25.0)
		self.assertEqual({c["campaign_name"]: c["replied"] for c in data["top_campaigns"]}["Mine Camp"], 1)
		frappe.get_doc({"doctype": "CRM Campaign Opt Out", "channel": "Email", "value": "stop.test@example.com", "source": "Manual", "lead": self.lead.name}).insert(ignore_permissions=True)
		self.assertEqual(self.nurturing()["totals"]["optouts"], 1)
		self.assertEqual(self.nurturing(user=self.rival.name)["totals"]["optouts"], 0)

	def test_empty_period_is_all_zeros(self):
		data = dashboard.get_nurturing("2001-01-01", "2001-01-03")
		self.assertEqual(data["totals"]["sent"], 0)
		self.assertEqual(len(data["by_day"]), 3)
		self.assertIsNone(data["totals"]["delivered_rate"])
		self.assertEqual(data["top_campaigns"], [])

	def test_nurtured_leads_are_counted_against_open_leads(self):
		data = self.nurturing(user=self.rep.name)
		self.assertEqual(data["nurtured"]["nurtured"], 2)
		self.assertGreaterEqual(data["nurtured"]["open_leads"], 2)


class TestDateBoundaries(FollowUpBase):
	"""Every report uses the same range: the days from the first to the last, both included, to the second.

	Records sit at 00:00:00 and 23:59:59 of the first and last day and just outside them (23:59:59 the day
	before, 00:00:00 the day after). The "previous period" is the same number of days that ends the day before.
	"""

	D = "2020-03-10"

	campaign = TestNurturing.campaign
	recipient = TestNurturing.recipient

	def setUp(self):
		super().setUp()
		self.rep = self.owner

	def test_sales_report_counts_calls_from_00_00_00_to_23_59_59(self):
		for when in ("2020-03-09 23:59:59", "2020-03-10 00:00:00", "2020-03-10 23:59:59", "2020-03-11 00:00:00", "2020-03-08 12:00:00"):
			self.log("Did Not Pick", followed_up_on=when)
		one = reports.get_sales_report(self.D, self.D)
		self.assertEqual(one["totals"]["calls"], 2)
		self.assertEqual(sum(r["calls"] for r in one["daily"]), 2)
		self.assertEqual(reports.get_sales_report("2020-03-09", self.D)["totals"]["calls"], 3)
		self.assertEqual(reports.get_sales_report(self.D, "2020-03-11")["totals"]["calls"], 3)
		# the previous period of one day is the day before; of two days (10 and 11) the 8th and the 9th
		self.assertEqual(reports.get_sales_report("2020-03-11", "2020-03-11")["previous"]["calls"], 2)
		self.assertEqual(reports.get_sales_report(self.D, "2020-03-11")["previous"]["calls"], 2)
		# the Excel / CSV call log is the same set of calls
		self.assertEqual(len(reports._call_log(self.D, self.D, None)), 2)
		self.assertEqual(len(reports._call_log("2020-03-09", "2020-03-11", None)), 4)

	def test_sales_report_counts_leads_created_on_the_boundary_days(self):
		extra = [
			frappe.get_doc({"doctype": "CRM Lead", "first_name": f"Edge{i}", "last_name": "Lead"}).insert(ignore_permissions=True)
			for i in range(3)
		]
		for lead, created in zip([self.lead, *extra], ("2020-03-09 23:59:59", "2020-03-10 00:00:00", "2020-03-10 23:59:59", "2020-03-11 00:00:00")):
			frappe.db.set_value("CRM Lead", lead.name, "creation", created, update_modified=False)
		self.assertEqual(reports.get_sales_report(self.D, self.D)["totals"]["leads"], 2)
		self.assertEqual(reports.get_sales_report("2020-03-11", "2020-03-11")["previous"]["leads"], 2)

	def test_follow_up_page_window_includes_the_whole_last_day(self):
		for when in ("2031-05-09 23:59:59", "2031-05-10 00:00:00", "2031-05-10 23:59:59", "2031-05-11 00:00:00"):
			self.log("Did Not Pick", next_follow_up_on=when)
		frappe.db.set_value("CRM Follow Up", {"reference_docname": self.lead.name}, "next_closed", 0)
		one = followups.get_workspace(tab="upcoming", from_date="2031-05-10", to_date="2031-05-10")
		self.assertEqual((one["total"], one["counts"]["upcoming"]), (2, 2))
		self.assertEqual(followups.get_workspace(tab="upcoming", from_date="2031-05-10", to_date="2031-05-11")["total"], 3)
		# reversed dates are put right, and a lone date is a one-day range
		self.assertEqual(followups.get_workspace(tab="upcoming", from_date="2031-05-11", to_date="2031-05-10")["total"], 3)
		self.assertEqual(followups.get_workspace(tab="upcoming", from_date="2031-05-10")["total"], 2)
		self.assertEqual(one["window"], {"from": "2031-05-10", "to": "2031-05-10"})

	def test_nurturing_counts_sends_from_00_00_00_to_23_59_59(self):
		camp = self.campaign("Edge Camp", self.rep.name)
		for when in ("2020-03-09 23:59:59", "2020-03-10 00:00:00", "2020-03-10 23:59:59", "2020-03-11 00:00:00", "2020-03-07 12:00:00"):
			self.recipient(camp, "Email", "Sent", when)
		one = dashboard.get_nurturing(self.D, self.D)
		self.assertEqual(one["totals"]["sent"], 2)
		self.assertEqual([(d["date"], d["Email"]) for d in one["by_day"]], [(self.D, 2)])
		# 10th to 12th is three days: the previous period is the 7th to the 9th
		three = dashboard.get_nurturing(self.D, "2020-03-12")
		self.assertEqual(three["totals"]["sent"], 3)
		self.assertEqual(three["previous"]["sent"], 2)
		self.assertEqual(len(three["by_day"]), 3)
		self.assertEqual(dashboard.get_nurturing("2020-03-11", "2020-03-11")["previous"]["sent"], 2)

	def test_meeting_counts_and_the_previous_period_of_the_crm_dashboard_card(self):
		fake = MeetingsTestCase_google()
		p1, p2 = fake
		with p1, p2, patch.object(frappe, "sendmail"):
			for start, end in (
				("2020-03-07 12:00:00", "2020-03-07 12:30:00"),
				("2020-03-09 23:30:00", "2020-03-09 23:59:59"),
				("2020-03-10 00:00:00", "2020-03-10 00:30:00"),
				("2020-03-10 23:30:00", "2020-03-10 23:59:59"),
				("2020-03-11 00:00:00", "2020-03-11 00:30:00"),
			):
				self.save(notify="none", add_video_conferencing=0, starts_on=start, ends_on=end)
		self.assertEqual(dashboard._count(self.D, self.D, None), 2)
		self.assertEqual(dashboard.get_meetings_in_period(self.D, self.D)["value"], 2)
		card = dashboard.get_meetings_in_period(self.D, "2020-03-12")
		# three days (10th to 12th): 3 meetings; the previous three days (7th to 9th): 2, so +50%
		self.assertEqual(card["value"], 3)
		self.assertEqual(round(card["delta"]), 50)

	def test_the_previous_period_is_as_long_as_the_current_one(self):
		# seven days, 10th to 16th: the calls of the 3rd to the 9th are the comparison
		for when in ("2020-03-02 23:59:59", "2020-03-03 00:00:00", "2020-03-09 23:59:59", "2020-03-10 00:00:00"):
			self.log("Interested", followed_up_on=when)
		data = reports.get_sales_report(self.D, "2020-03-16")
		self.assertEqual((data["totals"]["calls"], data["previous"]["calls"]), (1, 2))
		card = dashboard.get_call_connect_rate(self.D, "2020-03-16", None)
		self.assertEqual(card["value"], 100.0)
		self.assertEqual(card["delta"], 0.0)


def MeetingsTestCase_google():
	return google_patch(FakeGoogle())
