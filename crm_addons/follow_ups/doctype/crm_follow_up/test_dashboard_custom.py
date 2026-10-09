# Copyright (c) 2026, Coding Pro
"""A person's own dashboard layout, the Campaigns export, and leaving deals out for a team that does not use them.

Run on a bench:  bench --site <site> run-tests --app crm_addons --module crm_addons.follow_ups.doctype.crm_follow_up.test_dashboard_custom
"""

import io
import json

import frappe

from crm_addons import api, dashboard, reports
from crm_addons.meetings.doctype.crm_meeting.test_crm_meeting import TestCase as Base


class TestLayout(Base):
	def test_nothing_is_saved_until_someone_customises(self):
		self.assertIsNone(dashboard.get_layout("sales"))
		self.assertIsNone(dashboard.get_layout("campaigns"))

	def test_a_layout_is_saved_per_person_and_per_view_and_can_be_reset(self):
		layout = {"kpis": [{"id": "calls", "on": 1}, {"id": "stale", "on": 0}], "widgets": [{"id": "daily", "w": 6, "on": 1}, {"id": "team", "w": 3, "on": 0}]}
		saved = dashboard.save_layout("sales", json.dumps(layout))
		self.assertEqual(saved, {"kpis": layout["kpis"], "widgets": [{"id": "daily", "on": 1, "w": 6}, {"id": "team", "on": 0, "w": 3}]})
		self.assertEqual(dashboard.get_layout("sales"), saved)
		self.assertIsNone(dashboard.get_layout("campaigns"))  # the other view is untouched
		frappe.set_user(self.owner.name)
		self.assertIsNone(dashboard.get_layout("sales"))  # and so is somebody else
		frappe.set_user("Administrator")
		dashboard.reset_layout("sales")
		self.assertIsNone(dashboard.get_layout("sales"))

	def test_whatever_is_not_a_known_shape_is_dropped(self):
		junk = {"kpis": [{"id": "ok_one", "on": 1}, {"id": "has space", "on": 1}, {"id": "ok_one", "on": 0}, "text", {"id": "x" * 80}], "widgets": [{"id": "daily", "w": 5, "on": 1}, None]}
		out = dashboard.save_layout("sales", json.dumps(junk))
		self.assertEqual([k["id"] for k in out["kpis"]], ["ok_one"])
		self.assertEqual(out["widgets"], [{"id": "daily", "on": 1, "w": 3}])  # an odd width falls back to a half row
		with self.assertRaises(frappe.ValidationError):
			dashboard.save_layout("sales", json.dumps(["not", "a", "layout"]))
		with self.assertRaises(frappe.ValidationError):
			dashboard.save_layout("nope", "{}")

	def test_it_is_bounded(self):
		many = {"kpis": [{"id": f"k{i}", "on": 1} for i in range(500)], "widgets": []}
		self.assertLessEqual(len(dashboard.save_layout("sales", json.dumps(many))["kpis"]), dashboard.LAYOUT_MAX_ITEMS)

	def test_only_sales_roles_may_use_it(self):
		stranger = self.make_user("layout.stranger@example.org", "Website Manager")  # (made first: with the flag off, a site that can send mail commits when it creates a user)
		frappe.flags.in_test = False
		self.addCleanup(setattr, frappe.flags, "in_test", True)
		frappe.set_user(stranger.name)
		with self.assertRaises(frappe.PermissionError):
			dashboard.get_layout("sales")
		with self.assertRaises(frappe.PermissionError):
			dashboard.save_layout("sales", "{}")


class TestDealsAndCampaignExport(Base):
	def set_hide_deals(self, value):
		frappe.db.set_single_value("CRM Addons Settings", "hide_deals_menu", value)
		frappe.clear_document_cache("CRM Addons Settings", "CRM Addons Settings")

	def test_the_dashboards_know_whether_deals_are_in_use(self):
		self.set_hide_deals(1)
		self.assertEqual(api.get_client_config()["deals_enabled"], 0)
		self.set_hide_deals(0)
		self.assertEqual(api.get_client_config()["deals_enabled"], 1)

	def test_the_sales_workbook_leaves_deals_out_when_they_are_not_used(self):
		from openpyxl import load_workbook

		def sheets():
			frappe.response.clear()
			reports.export_report("2026-01-01", "2026-12-31", None, "xlsx")
			wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
			return wb, {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb}

		self.set_hide_deals(0)
		_wb, data = sheets()
		self.assertTrue(any("deals" in str(r[0]).lower() for r in data["Summary"]))
		self.assertIn("Converted", data["Team"][0])
		self.set_hide_deals(1)
		_wb, data = sheets()
		self.assertFalse(any("deals" in str(r[0]).lower() for r in data["Summary"]))
		self.assertNotIn("Converted", data["Team"][0])
		self.assertEqual(data["Sources"][0], ["Source", "Leads added"])
		self.assertNotIn("Converted to deal", [r[0] for r in data["Funnel"]])

	def test_the_campaigns_view_can_be_downloaded(self):
		import csv

		from openpyxl import load_workbook

		frappe.response.clear()
		reports.export_campaigns("2026-01-01", "2026-12-31", None, "xlsx")
		self.assertEqual(frappe.response["type"], "download")
		self.assertTrue(frappe.response["filename"].startswith("campaigns_2026-01-01_to_2026-12-31"))
		wb = load_workbook(io.BytesIO(frappe.response["filecontent"]))
		self.assertEqual(wb.sheetnames, ["Summary", "By day", "Channels", "Campaigns", "Status"])
		frappe.response.clear()
		reports.export_campaigns("2026-01-01", "2026-12-31", None, "csv")
		rows = list(csv.reader(io.StringIO(frappe.response["filecontent"].decode("utf-8-sig"))))
		self.assertEqual(rows[0][:3], ["Campaign", "Status", "Sent"])
		self.assertTrue(frappe.response["filename"].endswith(".csv"))

	def test_a_stranger_cannot_download_it(self):
		stranger = self.make_user("export.stranger@example.org", "Website Manager")
		frappe.flags.in_test = False
		self.addCleanup(setattr, frappe.flags, "in_test", True)
		frappe.set_user(stranger.name)
		with self.assertRaises(frappe.PermissionError):
			reports.export_campaigns()
